from copy import copy
from io import BytesIO
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter


def get_sheet_names_from_bytes(file_bytes):
    """Inspects an Excel file in memory and returns all worksheet names."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def parse_column_index(col_input):
    """Converts a column letter (e.g., 'B') or integer string (e.g., '2') to a 1-based integer."""
    if isinstance(col_input, int):
        return col_input
    col_str = str(col_input).strip()
    if col_str.isdigit():
        return int(col_str)
    return column_index_from_string(col_str)


def read_styles_from_order_file(file_bytes, start_row, style_col, sheet_name=None):
    """Reads style numbers from the uploaded Order Details file."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    col_idx = parse_column_index(style_col)
    styles = []

    for row in range(start_row, ws.max_row + 1):
        val = ws.cell(row=row, column=col_idx).value
        if val is not None:
            clean_val = str(val).strip()
            if clean_val:
                styles.append(clean_val)

    wb.close()
    # Deduplicate while preserving original order
    return list(dict.fromkeys(styles))


def detect_master_layout(ws, scan_limit=25):
    """
    Scans the master sheet to automatically locate:
    1. Header row (e.g., row 9)
    2. Style column
    3. Maximum table width across parameter rows (1-8) and header
    4. Image column
    """
    header_row = None
    style_col = None
    image_col = None

    style_target_keywords = ["STYLE #", "STYLE NO", "STYLE", "DESIGN #", "DESIGN NO"]

    # 1. Locate header row and style column
    for r in range(1, min(scan_limit, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                norm_val = str(val).strip().upper()
                if any(kw == norm_val or kw in norm_val for kw in style_target_keywords):
                    header_row = r
                    style_col = c
                    break
        if header_row is not None:
            break

    if header_row is None:
        raise ValueError("Could not automatically locate a header row with 'Style #' in the selected tab.")

    # 2. Determine table width across parameter rows & header, and locate image column
    max_col = style_col
    for r in range(1, header_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip() != "":
                max_col = max(max_col, c)
                if r == header_row:
                    norm_val = str(val).strip().upper()
                    if any(img_kw in norm_val for img_kw in ["IMAGE", "PHOTO", "PICTURE", "SKETCH"]):
                        image_col = c

    # Fallback for image column if not explicitly labeled (Column C is default 3)
    if image_col is None:
        image_col = 3

    return header_row, style_col, max_col, image_col


def get_image_row_col(img):
    """Returns the 1-indexed (row, col) anchor coordinates of an image."""
    if hasattr(img, "anchor"):
        anchor = img.anchor
        if hasattr(anchor, "_from"):
            return anchor._from.row + 1, anchor._from.col + 1
        elif isinstance(anchor, str):
            return coordinate_to_tuple(anchor)
    return None, None


def get_image_bytes(img):
    """Extracts raw binary data from an openpyxl Image object."""
    if callable(getattr(img, "_data", None)):
        return img._data()
    if hasattr(img, "ref"):
        if hasattr(img.ref, "read"):
            img.ref.seek(0)
            return img.ref.read()
    return None


def map_style_blocks(ws_b, header_row, style_col, max_col):
    """Identifies start and end rows for every style block in the master worksheet."""
    ignored_labels = {"TOTAL", "GRAND TOTAL", "SUBTOTAL", "STYLE #", "STYLE", "NONE", ""}
    style_starts = []

    for r in range(header_row + 1, ws_b.max_row + 1):
        val = ws_b.cell(row=r, column=style_col).value
        if val is not None:
            clean_val = str(val).strip()
            if clean_val.upper() not in ignored_labels:
                style_starts.append((clean_val, r))

    blocks = {}
    for i, (style, s_row) in enumerate(style_starts):
        if i < len(style_starts) - 1:
            e_row = style_starts[i + 1][1] - 1
        else:
            e_row = ws_b.max_row
            while e_row > s_row and all(
                ws_b.cell(row=e_row, column=c).value is None for c in range(1, max_col + 1)
            ):
                e_row -= 1

        blocks[style.upper()] = (style, s_row, e_row)

    return blocks


def copy_cell_range(ws_src, ws_dest, s_row, e_row, dest_start_row, max_col, serial_no=None):
    """
    Copies values/formulas, formatting, and row heights.
    - Translates relative formula row offsets while keeping absolute locks intact.
    - Overwrites Column A with serial number on the block's top row only when serial_no is passed.
    """
    for r in range(s_row, e_row + 1):
        cur_dest_row = dest_start_row + (r - s_row)
        row_offset = cur_dest_row - r

        if ws_src.row_dimensions[r].height is not None:
            ws_dest.row_dimensions[cur_dest_row].height = ws_src.row_dimensions[r].height

        for c in range(1, max_col + 1):
            src_cell = ws_src.cell(row=r, column=c)
            dest_cell = ws_dest.cell(row=cur_dest_row, column=c)

            # Assign serial number only for extracted data blocks (not for rows 1-9)
            if serial_no is not None and c == 1:
                if r == s_row:
                    dest_cell.value = serial_no
                else:
                    dest_cell.value = ""
            else:
                val = src_cell.value
                # Translate formulas to adapt relative row coordinates
                if isinstance(val, str) and val.startswith("="):
                    if row_offset != 0:
                        try:
                            dest_cell.value = Translator(
                                val, origin=src_cell.coordinate
                            ).translate_formula(row_offset=row_offset, col_offset=0)
                        except Exception:
                            dest_cell.value = val
                    else:
                        dest_cell.value = val
                else:
                    dest_cell.value = val

            if src_cell.has_style:
                dest_cell.font = copy(src_cell.font)
                dest_cell.border = copy(src_cell.border)
                dest_cell.fill = copy(src_cell.fill)
                dest_cell.number_format = copy(src_cell.number_format)
                dest_cell.protection = copy(src_cell.protection)
                dest_cell.alignment = copy(src_cell.alignment)


def copy_merged_cells(ws_src, ws_dest, s_row, e_row, dest_start_row):
    """Re-creates cell merges inside the extracted block."""
    row_offset = dest_start_row - s_row
    for rng in list(ws_src.merged_cells.ranges):
        if rng.min_row >= s_row and rng.max_row <= e_row:
            ws_dest.merge_cells(
                start_row=rng.min_row + row_offset,
                end_row=rng.max_row + row_offset,
                start_column=rng.min_col,
                end_column=rng.max_col,
            )


def add_fitted_image_to_block(ws_dest, img_bytes, dest_start_row, block_height, image_col=3):
    """Fits and anchors an image inside the target image column across up to 5 rows."""
    if not img_bytes:
        return

    new_img = Image(BytesIO(img_bytes))
    span_rows = min(5, block_height)
    start_row_idx = dest_start_row - 1
    end_row_idx = start_row_idx + span_rows

    marker_from = AnchorMarker(col=image_col - 1, colOff=150000, row=start_row_idx, rowOff=50000)
    marker_to = AnchorMarker(col=image_col, colOff=-150000, row=end_row_idx, rowOff=-50000)

    new_img.anchor = TwoCellAnchor(_from=marker_from, to=marker_to, editAs="oneCell")
    ws_dest.add_image(new_img)


def process_workbook_extraction(order_bytes, master_bytes, master_sheet_name, order_start_row, order_style_col):
    """Runs extraction entirely in memory and returns (output_excel_bytes, found_styles, missing_styles)."""
    # 1. Read style list from Order File
    target_styles = read_styles_from_order_file(order_bytes, order_start_row, order_style_col)

    # 2. Load Master File with formulas preserved
    wb_b = openpyxl.load_workbook(BytesIO(master_bytes), data_only=False)
    if master_sheet_name not in wb_b.sheetnames:
        raise ValueError(f"Sheet '{master_sheet_name}' not found in the master file.")
    ws_b = wb_b[master_sheet_name]

    # 3. Detect master structure dynamically
    header_row, style_col, max_col, image_col = detect_master_layout(ws_b)
    style_blocks = map_style_blocks(ws_b, header_row, style_col, max_col)

    # 4. Prepare target workbook
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = f"Master_{master_sheet_name}"

    # Copy top parameter table (rows 1–8) and main table header (header_row) directly
    copy_cell_range(ws_b, ws_out, 1, header_row, 1, max_col, serial_no=None)
    copy_merged_cells(ws_b, ws_out, 1, header_row, 1)

    # Copy column widths
    for c in range(1, max_col + 1):
        col_letter = get_column_letter(c)
        w = ws_b.column_dimensions[col_letter].width
        if w:
            ws_out.column_dimensions[col_letter].width = w

    # Style rows start directly beneath the main header row (row 10 when header is row 9)
    dest_current_row = header_row + 1
    serial_number = 1
    found_styles = []
    missing_styles = []

    # 5. Extract style blocks
    for style in target_styles:
        key = style.upper()
        if key in style_blocks:
            _, s_row, e_row = style_blocks[key]
            block_height = (e_row - s_row) + 1

            # Extract image bytes for this block if available
            block_img_bytes = None
            if hasattr(ws_b, "_images"):
                for img in ws_b._images:
                    img_row, _ = get_image_row_col(img)
                    if img_row and s_row <= img_row <= e_row:
                        block_img_bytes = get_image_bytes(img)
                        break

            # Copy contents & merge ranges with formula row-offset translation
            copy_cell_range(ws_b, ws_out, s_row, e_row, dest_current_row, max_col, serial_no=serial_number)
            copy_merged_cells(ws_b, ws_out, s_row, e_row, dest_current_row)

            # Re-insert image
            if block_img_bytes:
                add_fitted_image_to_block(ws_out, block_img_bytes, dest_current_row, block_height, image_col)

            dest_current_row += block_height
            serial_number += 1
            found_styles.append(style)
        else:
            missing_styles.append(style)

    wb_b.close()

    # Save output to buffer
    out_buffer = BytesIO()
    wb_out.save(out_buffer)
    out_buffer.seek(0)

    return out_buffer.getvalue(), found_styles, missing_styles