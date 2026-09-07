from copy import copy
from io import BytesIO
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils import (
    column_index_from_string,
    coordinate_to_tuple,
    get_column_letter,
)

# ==============================================================================
# CONFIGURATION
# ==============================================================================
FILE_A_PATH = "Order Details 28.08.2026.xlsx"  # Path to File A
FILE_B_PATH = "ANRO Master DCB Format.xlsx"  # Path to File B
OUTPUT_PATH = "Extracted_Master_Data.xlsx"  # Destination output path

# File A Settings
SHEET_A_NAME = "Sheet1"
STYLE_COL_A = 2  # Column B
DATA_START_ROW_A = 6  # First data row after header (row 5)

# File B Settings
SHEET_B_NAME = "14K SI"
HEADER_ROW_B = 9  # Table header row (Sr#, Style #, Image, etc.)
STYLE_COL_B = 2  # Column B
MAX_COL_NAME = "AI"  # Columns A to AI
MAX_COL_B = column_index_from_string(MAX_COL_NAME)  # 35

DEDUPLICATE_STYLES = True
# ==============================================================================


def get_image_row_col(img):
  """Returns the 1-indexed (row, col) where an image is anchored."""
  if hasattr(img, "anchor"):
    anchor = img.anchor
    if hasattr(anchor, "_from"):
      return anchor._from.row + 1, anchor._from.col + 1
    elif isinstance(anchor, str):
      return coordinate_to_tuple(anchor)
  return None, None


def get_image_bytes(img):
  """Safely extracts raw binary bytes from an openpyxl Image object."""
  if callable(getattr(img, "_data", None)):
    return img._data()
  if hasattr(img, "ref"):
    if hasattr(img.ref, "read"):
      img.ref.seek(0)
      return img.ref.read()
  return None


def read_styles_from_file_a(file_path, sheet_name, start_row, style_col):
  """Reads style numbers row-by-row from Column B in File A."""
  wb = openpyxl.load_workbook(file_path, data_only=True)
  ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

  styles = []
  for row in range(start_row, ws.max_row + 1):
    val = ws.cell(row=row, column=style_col).value
    if val is not None:
      clean_val = str(val).strip()
      if clean_val:
        styles.append(clean_val)

  wb.close()
  if DEDUPLICATE_STYLES:
    return list(dict.fromkeys(styles))
  return styles


def map_style_blocks(ws_b, header_row, style_col):
  """Identifies start and end row for each style block in File B."""
  ignored_labels = {
      "TOTAL",
      "GRAND TOTAL",
      "SUBTOTAL",
      "STYLE #",
      "STYLE",
      "NONE",
      "",
  }
  style_starts = []

  for r in range(header_row + 1, ws_b.max_row + 1):
    val = ws_b.cell(row=r, column=style_col).value
    if val is not None:
      clean_val = str(val).strip()
      if clean_val.upper() not in ignored_labels:
        style_starts.append((clean_val, r))

  blocks = {}
  for i, (style, s_row) in enumerate(style_starts):
    if i < len(style_starts) - 1:
      e_row = style_starts[i + 1][1] - 1
    else:
      e_row = ws_b.max_row
      while e_row > s_row and all(
          ws_b.cell(row=e_row, column=c).value is None
          for c in range(1, MAX_COL_B + 1)
      ):
        e_row -= 1

    blocks[style.upper()] = (style, s_row, e_row)

  return blocks


def copy_cell_range(
    ws_src, ws_dest, s_row, e_row, dest_start_row, max_col, serial_no=None
):
  """Copies cell values, styles, and row dimensions. Overwrites Column A with serial_no."""
  for r in range(s_row, e_row + 1):
    cur_dest_row = dest_start_row + (r - s_row)

    if ws_src.row_dimensions[r].height is not None:
      ws_dest.row_dimensions[cur_dest_row].height = ws_src.row_dimensions[
          r
      ].height

    for c in range(1, max_col + 1):
      src_cell = ws_src.cell(row=r, column=c)
      dest_cell = ws_dest.cell(row=cur_dest_row, column=c)

      # If this is Column A (Sr.#) and it's the first row of the style block, write serial number
      if c == 1 and r == s_row and serial_no is not None:
        dest_cell.value = serial_no
      elif c == 1 and r > s_row:
        dest_cell.value = ""  # Leave subsequent rows of block blank for Col A
      else:
        dest_cell.value = src_cell.value

      if src_cell.has_style:
        dest_cell.font = copy(src_cell.font)
        dest_cell.border = copy(src_cell.border)
        dest_cell.fill = copy(src_cell.fill)
        dest_cell.number_format = copy(src_cell.number_format)
        dest_cell.protection = copy(src_cell.protection)
        dest_cell.alignment = copy(src_cell.alignment)


def copy_merged_cells(ws_src, ws_dest, s_row, e_row, dest_start_row):
  """Re-creates merged cells within the copied block."""
  row_offset = dest_start_row - s_row
  for rng in list(ws_src.merged_cells.ranges):
    if rng.min_row >= s_row and rng.max_row <= e_row:
      ws_dest.merge_cells(
          start_row=rng.min_row + row_offset,
          end_row=rng.max_row + row_offset,
          start_column=rng.min_col,
          end_column=rng.max_col,
      )


def add_fitted_image_to_block(
    ws_dest, img_bytes, dest_start_row, block_height
):
  """Anchors the image across Column C and up to 5 rows high using TwoCellAnchor."""
  if not img_bytes:
    return

  new_img = Image(BytesIO(img_bytes))

  # 0-indexed column 2 = Column C, column 3 = Column D boundary
  # 0-indexed row start = dest_start_row - 1
  # Height span = 5 rows (or total block height if smaller)
  span_rows = min(5, block_height)

  start_row_idx = dest_start_row - 1
  end_row_idx = start_row_idx + span_rows

  marker_from = AnchorMarker(col=2, colOff=150000, row=start_row_idx, rowOff=50000)
  marker_to = AnchorMarker(col=3, colOff=-150000, row=end_row_idx, rowOff=-50000)

  new_img.anchor = TwoCellAnchor(
      _from=marker_from, to=marker_to, editAs="oneCell"
  )
  ws_dest.add_image(new_img)


def main():
  print("1. Reading styles from File A...")
  target_styles = read_styles_from_file_a(
      FILE_A_PATH, SHEET_A_NAME, DATA_START_ROW_A, STYLE_COL_A
  )
  print(f"   Found {len(target_styles)} styles in File A to search.")

  print("2. Loading File B...")
  wb_b = openpyxl.load_workbook(FILE_B_PATH, data_only=True)

  sheets_to_search = (
      [wb_b[SHEET_B_NAME]] if SHEET_B_NAME in wb_b.sheetnames else wb_b.worksheets
  )

  style_index = {}
  for ws in sheets_to_search:
    blocks = map_style_blocks(ws, HEADER_ROW_B, STYLE_COL_B)
    for k, v in blocks.items():
      style_index[k] = (ws, v)

  wb_out = openpyxl.Workbook()
  ws_out = wb_out.active
  ws_out.title = "Extracted Master Styles"

  src_header_ws = sheets_to_search[0]
  copy_cell_range(src_header_ws, ws_out, HEADER_ROW_B, HEADER_ROW_B, 1, MAX_COL_B)
  copy_merged_cells(src_header_ws, ws_out, HEADER_ROW_B, HEADER_ROW_B, 1)

  for c in range(1, MAX_COL_B + 1):
    col_letter = get_column_letter(c)
    w = src_header_ws.column_dimensions[col_letter].width
    if w:
      ws_out.column_dimensions[col_letter].width = w

  dest_current_row = 2
  extracted_count = 0
  serial_number = 1

  print("3. Extracting style blocks...")
  for style in target_styles:
    key = style.upper()
    if key in style_index:
      ws_src, (_, s_row, e_row) = style_index[key]
      block_height = (e_row - s_row) + 1

      # Find image bytes from source block if present
      block_image_bytes = None
      if hasattr(ws_src, "_images"):
        for img in ws_src._images:
          img_row, _ = get_image_row_col(img)
          if img_row and s_row <= img_row <= e_row:
            block_image_bytes = get_image_bytes(img)
            break

      # Copy cells, rows, and apply serial number
      copy_cell_range(
          ws_src,
          ws_out,
          s_row,
          e_row,
          dest_current_row,
          MAX_COL_B,
          serial_no=serial_number,
      )
      copy_merged_cells(ws_src, ws_out, s_row, e_row, dest_current_row)

      # Add fitted image across Column C and up to 5 rows
      if block_image_bytes:
        add_fitted_image_to_block(
            ws_out, block_image_bytes, dest_current_row, block_height
        )

      dest_current_row += block_height
      extracted_count += 1
      serial_number += 1
      print(f"   [MATCH] {style} extracted with Sr.# {serial_number - 1}")
    else:
      print(f"   [NOT FOUND] Style '{style}' not present in File B.")

  wb_out.save(OUTPUT_PATH)
  wb_b.close()
  print(
      f"\nDone! Successfully extracted {extracted_count} style blocks to"
      f" '{OUTPUT_PATH}'."
  )


if __name__ == "__main__":
  main()